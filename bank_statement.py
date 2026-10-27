"""Safe, local parsing of user-uploaded CSV/XLSX bank statements."""
import csv
import io
import re
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation

MAX_UPLOAD_BYTES = 8 * 1024 * 1024
MAX_ROWS = 2000

DATE_HEADERS = {"date", "transaction date", "booking date", "value date", "дата", "дата операции", "дата транзакции", "дата проводки", "buchungstag", "datum", "fecha", "data", "data operacji"}
DESC_HEADERS = {"description", "details", "transaction", "merchant", "payee", "name", "memo", "purpose", "описание", "назначение", "детали", "операция", "платеж", "платёж", "получатель", "отправитель", "контрагент", "опис операції", "призначення", "деталі", "merchant name", "verwendungszweck", "beschreibung", "libellé", "concepto", "descrizione", "tytuł"}
AMOUNT_HEADERS = {"amount", "sum", "value", "transaction amount", "сумма", "сума", "сумма операции", "сума операції", "betrag", "importe", "kwota", "montant", "importo"}
DEBIT_HEADERS = {"debit", "withdrawal", "outgoing", "money out", "расход", "списание", "списано", "дебет", "витрата", "списання", "Soll"}
CREDIT_HEADERS = {"credit", "deposit", "incoming", "money in", "income", "поступление", "зачисление", "доход", "кредит", "надходження", "зарахування", "Haben"}
CURRENCY_HEADERS = {"currency", "ccy", "валюта", "währung", "devise", "moneda", "waluta"}
TYPE_HEADERS = {"type", "transaction type", "тип", "тип операции", "тип операції", "art"}
TRANSFER_WORDS = ("transfer", "перевод между", "переказ між", "own account", "internal transfer", "перевод на свой", "переказ на власний")


def _norm(value):
    return re.sub(r"\s+", " ", str(value or "").strip().lower().replace("\ufeff", ""))


def _parse_amount(value):
    if value is None:
        return None
    if isinstance(value, (int, float, Decimal)):
        try:
            result = Decimal(str(value))
            return result if result.is_finite() else None
        except InvalidOperation:
            return None
    s = str(value).strip().replace("\u00a0", "").replace(" ", "").replace("'", "").replace("’", "")
    if not s or s in {"-", "—", "–"}:
        return None
    negative = s.startswith("(") and s.endswith(")")
    s = s.strip("()")
    s = re.sub(r"[^0-9,\.\-+]", "", s)
    if not s or s in {"-", "+", ".", ","}:
        return None
    if "," in s and "." in s:
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    elif "," in s:
        tail = s.rsplit(",", 1)[-1]
        s = s.replace(",", "." if len(tail) in (1, 2) else "")
    elif s.count(".") > 1:
        parts = s.split(".")
        s = "".join(parts[:-1]) + "." + parts[-1] if len(parts[-1]) in (1, 2) else "".join(parts)
    try:
        n = Decimal(s)
        if negative:
            n = -abs(n)
        return n if n.is_finite() else None
    except InvalidOperation:
        return None


def _parse_date(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)) and 20000 <= value <= 80000:
        try:
            return (datetime(1899, 12, 30) + timedelta(days=float(value))).date().isoformat()
        except (OverflowError, ValueError):
            return None
    s = str(value or "").strip()
    if not s:
        return None
    for fmt in ("%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%d.%m.%Y", "%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y", "%Y/%m/%d", "%d.%m.%y"):
        try:
            return datetime.strptime(s[:19], fmt).date().isoformat()
        except ValueError:
            pass
    try:
        return datetime.fromisoformat(s.replace("Z", "+00:00")).date().isoformat()
    except ValueError:
        return None


def _header_index(headers, choices):
    for i, h in enumerate(headers):
        hn = _norm(h)
        if hn in choices:
            return i
    for i, h in enumerate(headers):
        hn = _norm(h)
        for candidate in choices:
            if len(candidate) < 4 or candidate not in hn:
                continue
            # Avoid mistaking a "transaction type" column for a description.
            if candidate == "transaction" and "type" in hn:
                continue
            return i
    return None


def _decode_csv(data):
    for encoding in ("utf-8-sig", "utf-8", "cp1251", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def parse_bank_statement(filename, data):
    if len(data) > MAX_UPLOAD_BYTES:
        raise ValueError("Файл больше 8 МБ. Раздели выписку на несколько файлов.")
    ext = str(filename or "").lower().rsplit(".", 1)[-1]
    if ext not in {"csv", "xlsx"}:
        raise ValueError("Пока поддерживаются CSV и XLSX. PDF добавим отдельным этапом.")
    if ext == "xlsx":
        try:
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            ws = wb.active
            raw = list(ws.iter_rows(values_only=True))
            wb.close()
        except Exception as exc:
            raise ValueError("Не удалось прочитать XLSX. Сохрани выписку заново в Excel.") from exc
    else:
        decoded = _decode_csv(data)
        try:
            dialect = csv.Sniffer().sniff(decoded[:8192], delimiters=",;\t|")
            delimiter = dialect.delimiter
        except csv.Error:
            delimiter = ";" if decoded.count(";") > decoded.count(",") else ","
        raw = list(csv.reader(io.StringIO(decoded), delimiter=delimiter))
    if not raw:
        raise ValueError("Файл пустой.")
    raw = [list(row) for row in raw if row and any(str(v or "").strip() for v in row)]
    if not raw:
        raise ValueError("В файле нет строк с данными.")
    header_pos = None
    indexes = None
    for pos, row in enumerate(raw[:15]):
        headers = [_norm(v) for v in row]
        date_i = _header_index(headers, DATE_HEADERS)
        desc_i = _header_index(headers, DESC_HEADERS)
        amount_i = _header_index(headers, AMOUNT_HEADERS)
        debit_i = _header_index(headers, DEBIT_HEADERS)
        credit_i = _header_index(headers, CREDIT_HEADERS)
        if date_i is not None and (desc_i is not None or amount_i is not None or debit_i is not None or credit_i is not None):
            header_pos = pos
            indexes = (headers, date_i, desc_i, amount_i, debit_i, credit_i,
                       _header_index(headers, CURRENCY_HEADERS), _header_index(headers, TYPE_HEADERS))
            break
    if header_pos is None:
        raise ValueError("Не распознаны столбцы выписки. Нужны дата, описание операции и сумма (или отдельные столбцы дебет/кредит).")
    headers, date_i, desc_i, amount_i, debit_i, credit_i, currency_i, type_i = indexes
    rows = []
    skipped = 0
    for raw_row in raw[header_pos + 1:]:
        if len(rows) >= MAX_ROWS:
            break
        row = list(raw_row)
        def cell(i):
            return row[i] if i is not None and i < len(row) else None
        tx_date = _parse_date(cell(date_i))
        desc = str(cell(desc_i) or "").strip()[:180] if desc_i is not None else ""
        debit = _parse_amount(cell(debit_i)) if debit_i is not None else None
        credit = _parse_amount(cell(credit_i)) if credit_i is not None else None
        if debit_i is not None or credit_i is not None:
            if (debit is None or debit == 0) and (credit is None or credit == 0):
                amount = _parse_amount(cell(amount_i)) if amount_i is not None else None
            elif debit not in (None, Decimal("0")):
                amount = -abs(debit)
            else:
                amount = abs(credit or Decimal("0"))
        else:
            amount = _parse_amount(cell(amount_i)) if amount_i is not None else None
        if not tx_date or amount is None or amount == 0:
            skipped += 1
            continue
        description_lower = desc.lower()
        raw_type = _norm(cell(type_i)) if type_i is not None else ""
        is_transfer = any(word in description_lower for word in TRANSFER_WORDS) or "transfer" in raw_type or "переказ" in raw_type or "перевод" in raw_type
        kind = "transfer" if is_transfer else ("income" if amount > 0 else "expense")
        rows.append({
            "date": tx_date,
            "title": desc or ("Доход из выписки" if amount > 0 else "Расход из выписки"),
            "amount": float(abs(amount).quantize(Decimal("0.01"))),
            "type": kind,
            "currency": str(cell(currency_i) or "").strip().upper()[:8] if currency_i is not None else "",
        })
    if not rows:
        raise ValueError("Не удалось найти операции. Проверь формат даты, сумму и названия столбцов.")
    return {"rows": rows, "skipped": skipped, "truncated": len(raw) - header_pos - 1 > MAX_ROWS}
